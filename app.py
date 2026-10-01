import os
import re
import tempfile
from pathlib import Path
from io import BytesIO
from typing import Any

import faiss
import numpy as np
import streamlit as st
from docx import Document
from groq import Groq
from pypdf import PdfReader
from rank_bm25 import BM25Okapi
from sentence_transformers import SentenceTransformer

APP_TITLE = "AI Document Assistant"
MODEL_NAME = "all-MiniLM-L6-v2"
DEFAULT_GROQ_MODEL = "openai/gpt-oss-120b"
CHUNK_SIZE = 1000
CHUNK_OVERLAP = 200
TOP_K = 5

st.set_page_config(page_title=APP_TITLE, page_icon="📚", layout="wide")

@st.cache_resource(show_spinner="Loading embedding model...")
def load_embedding_model():
    return SentenceTransformer(MODEL_NAME)

def extract_pdf(file_bytes: bytes, filename: str) -> list[dict[str, Any]]:
    reader = PdfReader(BytesIO(file_bytes))
    sections = []
    for page_num, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            sections.append({"text": text, "filename": filename, "page": page_num})
    return sections

def extract_docx(file_bytes: bytes, filename: str) -> list[dict[str, Any]]:
    doc = Document(BytesIO(file_bytes))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    text = "\n".join(parts).strip()
    return [{"text": text, "filename": filename, "page": None}] if text else []

def extract_text(file_bytes: bytes, filename: str) -> list[dict[str, Any]]:
    text = file_bytes.decode("utf-8-sig", errors="replace").strip()
    return [{"text": text, "filename": filename, "page": None}] if text else []

def extract_document(file_bytes: bytes, filename: str) -> list[dict[str, Any]]:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return extract_pdf(file_bytes, filename)
    if ext == ".docx":
        return extract_docx(file_bytes, filename)
    if ext in {".txt", ".md"}:
        return extract_text(file_bytes, filename)
    raise ValueError(f"Unsupported file type: {ext or 'unknown'}")

def split_into_chunks(text: str, chunk_size: int = CHUNK_SIZE,
                      overlap: int = CHUNK_OVERLAP) -> list[str]:
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return []
    if chunk_size <= 0 or overlap < 0 or overlap >= chunk_size:
        raise ValueError("Chunk size must be positive and overlap must be between 0 and chunk size.")
    chunks = []
    start = 0
    while start < len(text):
        end = min(start + chunk_size, len(text))
        # Prefer a nearby word boundary without creating tiny chunks.
        if end < len(text):
            boundary = text.rfind(" ", start + int(chunk_size * 0.7), end)
            if boundary > start:
                end = boundary
        piece = text[start:end].strip()
        if piece:
            chunks.append(piece)
        if end >= len(text):
            break
        start = max(start + 1, end - overlap)
    return chunks

def build_chunks(files, chunk_size: int, overlap: int) -> list[dict[str, Any]]:
    result = []
    for uploaded in files:
        filename = uploaded.name
        sections = extract_document(uploaded.getvalue(), filename)
        for section in sections:
            for idx, piece in enumerate(split_into_chunks(section["text"], chunk_size, overlap), 1):
                result.append({
                    "text": piece,
                    "filename": filename,
                    "page": section["page"],
                    "chunk_index": idx,
                })
    return result

def tokenize(text: str) -> list[str]:
    return re.findall(r"\b[\w'-]+\b", text.lower())

def build_search_index(chunks: list[dict[str, Any]], model):
    if not chunks:
        raise ValueError("No readable text was found in the uploaded documents.")
    texts = [c["text"] for c in chunks]
    vectors = model.encode(
        texts, convert_to_numpy=True, normalize_embeddings=True,
        show_progress_bar=False
    ).astype("float32")
    index = faiss.IndexFlatIP(vectors.shape[1])
    index.add(vectors)
    tokenized = [tokenize(t) for t in texts]
    bm25 = BM25Okapi(tokenized)
    return index, vectors, bm25

def hybrid_search(query: str, chunks: list[dict[str, Any]], model,
                  index, bm25, top_k: int = TOP_K,
                  semantic_weight: float = 0.65):
    query_vector = model.encode(
        [query], convert_to_numpy=True, normalize_embeddings=True
    ).astype("float32")
    semantic_scores, semantic_ids = index.search(query_vector, min(len(chunks), len(chunks)))
    semantic = {int(i): float(s) for s, i in zip(semantic_scores[0], semantic_ids[0]) if i >= 0}

    raw_bm25 = np.asarray(bm25.get_scores(tokenize(query)), dtype=np.float32)
    if raw_bm25.size and raw_bm25.max() > raw_bm25.min():
        keyword = (raw_bm25 - raw_bm25.min()) / (raw_bm25.max() - raw_bm25.min())
    else:
        keyword = np.zeros(len(chunks), dtype=np.float32)

    # Cosine similarity can be negative; map to 0..1 before combining.
    sem = np.zeros(len(chunks), dtype=np.float32)
    for i, score in semantic.items():
        sem[i] = (score + 1.0) / 2.0
    combined = semantic_weight * sem + (1.0 - semantic_weight) * keyword
    ranked = np.argsort(combined)[::-1][:min(top_k, len(chunks))]
    return [
        {**chunks[int(i)], "score": float(combined[int(i)]),
         "semantic_score": float(sem[int(i)]), "keyword_score": float(keyword[int(i)])}
        for i in ranked
    ]

def get_groq_client(api_key: str) -> Groq:
    return Groq(api_key=api_key)

def answer_question(query: str, results: list[dict[str, Any]], api_key: str,
                    model_name: str) -> str:
    context_parts = []
    for i, item in enumerate(results, start=1):
        location = item["filename"]
        if item["page"] is not None:
            location += f", page {item['page']}"
        context_parts.append(f"[Source {i}: {location}]\n{item['text']}")
    context = "\n\n".join(context_parts)
    client = get_groq_client(api_key)
    completion = client.chat.completions.create(
        model=model_name,
        temperature=0.2,
        messages=[
            {"role": "system", "content":
             "You are a careful document question-answering assistant. "
             "Answer using only the supplied document context. If the context does not "
             "contain the answer, say you could not find it in the uploaded documents. "
             "Cite supporting sources using [Source 1], [Source 2], etc. Do not invent facts."},
            {"role": "user", "content": f"DOCUMENT CONTEXT:\n{context}\n\nQUESTION:\n{query}"}
        ],
    )
    return completion.choices[0].message.content or "No answer was returned."

def extract_drive_id(value: str) -> str:
    value = value.strip()
    match = re.search(r"(?:/d/|id=)([A-Za-z0-9_-]+)", value)
    if match:
        return match.group(1)
    if re.fullmatch(r"[A-Za-z0-9_-]{15,}", value):
        return value
    raise ValueError("Enter a public Google Drive file URL or file ID.")

def download_drive_file(value: str) -> tuple[bytes, str]:
    import gdown
    file_id = extract_drive_id(value)
    url = f"https://drive.google.com/uc?id={file_id}"
    with tempfile.TemporaryDirectory() as temp_dir:
        output = os.path.join(temp_dir, "drive_download")
        result = gdown.download(url, output, quiet=True, fuzzy=True)
        if not result or not os.path.exists(output):
            raise ValueError("Download failed. Make sure the file is shared as 'Anyone with the link'.")
        data = Path(output).read_bytes()
    # File extension may be inferred from the uploaded Drive URL only imperfectly.
    name_match = re.search(r"/([^/?#]+)\.(pdf|docx|txt|md)(?:[?#]|$)", value, re.I)
    filename = name_match.group(0).split("/")[-1].split("?")[0] if name_match else "google_drive_document.pdf"
    return data, filename

def reset_index():
    for key in ("chunks", "faiss_index", "embeddings", "bm25", "indexed_files"):
        st.session_state.pop(key, None)

st.title("📚 AI Document Assistant")
st.caption("Upload documents or add a public Google Drive file, then ask questions grounded in their content.")

with st.sidebar:
    st.header("Settings")
    chunk_size = st.slider("Chunk size (characters)", 300, 2000, CHUNK_SIZE, 100)
    overlap = st.slider("Chunk overlap (characters)", 0, min(500, chunk_size - 1), min(CHUNK_OVERLAP, chunk_size - 1), 50)
    top_k = st.slider("Retrieved chunks", 1, 10, TOP_K)
    semantic_weight = st.slider("Semantic search weight", 0.0, 1.0, 0.65, 0.05)
    groq_model = st.text_input("Groq model", value=DEFAULT_GROQ_MODEL)
    st.caption("The embedding model downloads on first use.")

api_key = st.secrets.get("GROQ_API_KEY", os.environ.get("GROQ_API_KEY", ""))
with st.sidebar:
    if api_key:
        st.success("Groq API key detected")
    else:
        st.warning("Add GROQ_API_KEY in Streamlit secrets or environment variables.")

tab_upload, tab_drive, tab_chat = st.tabs(["📁 Upload documents", "🔗 Google Drive", "💬 Ask questions"])

with tab_upload:
    uploaded_files = st.file_uploader(
        "Choose PDF, DOCX, TXT, or Markdown files",
        type=["pdf", "docx", "txt", "md"],
        accept_multiple_files=True,
        key="local_uploads"
    )
    if st.button("Process uploaded documents", type="primary", disabled=not uploaded_files):
        try:
            with st.spinner("Extracting text, splitting chunks, and building search index..."):
                chunks = build_chunks(uploaded_files, chunk_size, overlap)
                model = load_embedding_model()
                index, vectors, bm25 = build_search_index(chunks, model)
                st.session_state["chunks"] = chunks
                st.session_state["faiss_index"] = index
                st.session_state["embeddings"] = vectors
                st.session_state["bm25"] = bm25
                st.session_state["indexed_files"] = [f.name for f in uploaded_files]
            st.success(f"Indexed {len(uploaded_files)} file(s) into {len(chunks)} chunks.")
        except Exception as exc:
            st.error(f"Could not process documents: {exc}")

with tab_drive:
    st.write("Add a publicly shared Google Drive file. The file must be shared with 'Anyone with the link'.")
    drive_url = st.text_input("Google Drive file URL or file ID")
    if st.button("Download and index Drive file", disabled=not drive_url.strip()):
        try:
            with st.spinner("Downloading and indexing file..."):
                data, filename = download_drive_file(drive_url)
                class MemoryUpload:
                    name = filename
                    def getvalue(self):
                        return data
                chunks = build_chunks([MemoryUpload()], chunk_size, overlap)
                model = load_embedding_model()
                index, vectors, bm25 = build_search_index(chunks, model)
                st.session_state["chunks"] = chunks
                st.session_state["faiss_index"] = index
                st.session_state["embeddings"] = vectors
                st.session_state["bm25"] = bm25
                st.session_state["indexed_files"] = [filename]
            st.success(f"Indexed Google Drive file into {len(chunks)} chunks.")
        except Exception as exc:
            st.error(f"Could not process Google Drive file: {exc}")

with tab_chat:
    if "chunks" not in st.session_state:
        st.info("Upload and process documents first.")
    else:
        st.success(f"Ready: {len(st.session_state['chunks'])} searchable chunks")
        question = st.text_area("Ask a question about your documents", placeholder="e.g., What are the main findings?")
        if st.button("Get answer", type="primary", disabled=not question.strip()):
            if not api_key:
                st.error("Groq API key is missing. Add it in Streamlit secrets or the environment.")
            else:
                try:
                    with st.spinner("Searching documents and generating an answer..."):
                        model = load_embedding_model()
                        results = hybrid_search(
                            question, st.session_state["chunks"], model,
                            st.session_state["faiss_index"], st.session_state["bm25"],
                            top_k=top_k, semantic_weight=semantic_weight
                        )
                        answer = answer_question(question, results, api_key, groq_model)
                    st.subheader("Answer")
                    st.markdown(answer)
                    st.subheader("Sources")
                    for i, item in enumerate(results, start=1):
                        location = item["filename"]
                        if item["page"] is not None:
                            location += f" · Page {item['page']}"
                        with st.expander(f"[Source {i}] {location} · score {item['score']:.3f}"):
                            st.write(item["text"])
                except Exception as exc:
                    st.error(f"Question answering failed: {exc}")

with st.sidebar:
    st.divider()
    if "chunks" in st.session_state:
        st.caption(f"Indexed files: {', '.join(st.session_state.get('indexed_files', []))}")
        if st.button("Clear indexed documents"):
            reset_index()
            st.rerun()
